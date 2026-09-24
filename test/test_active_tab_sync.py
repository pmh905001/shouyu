import time

import openpyxl

from shouyu.service.excel import KbExcel
from shouyu.service.plan import BACKLOG_SHEET_LIFE, BACKLOG_SHEET_WORK, PLAN_HEADER_TEXT


def _assert_active_tab_consistent(wb: openpyxl.Workbook, today: str) -> None:
    today_ws = wb[today]
    assert wb.active.title == today
    assert today_ws.views.sheetView[0].tabSelected is True
    assert wb.views[0].activeTab == wb.sheetnames.index(today)
    assert wb.views[0].firstSheet == wb.sheetnames.index(today)
    for name in wb.sheetnames:
        if name == today:
            continue
        assert wb[name].views.sheetView[0].tabSelected is False


def test_active_tab_sync_after_pinned_sheet_regroup(tmp_path):
    """move_sheet() must not leave activeTab and tabSelected pointing at
    different sheets — that triggers Google Sheets' active-tab prompt."""
    path = tmp_path / "daily.xlsx"
    today = time.strftime("%Y-%m-%d")

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    day_ws = wb.create_sheet(today)
    day_ws["A1"] = PLAN_HEADER_TEXT
    backlog_work = wb.create_sheet(BACKLOG_SHEET_WORK)
    backlog_work["A1"] = "backlog"
    backlog_life = wb.create_sheet(BACKLOG_SHEET_LIFE)
    backlog_life["A1"] = "backlog"
    # Pin pinned sheets in the wrong order so regroup will move them.
    wb._sheets = [day_ws, backlog_life, backlog_work]
    wb.save(path)

    kb = KbExcel(str(path))
    kb.active_worksheet
    kb.mark_changed()
    kb.force_save()

    with openpyxl.load_workbook(path) as saved:
        _assert_active_tab_consistent(saved, today)
